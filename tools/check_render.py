"""Fail if markdown inline markers leaked into a rendered DOCX.

Why this exists as a gate rather than a one-off check
------------------------------------------------------
`find_marker_issues.py` inspects the SOURCE and reports lines with an odd number of `**`.
That is necessary but not sufficient: 1.1.0's source has 36 such lines (hard-wrapped bold
spanning a line break) and its DOCX renders all of them correctly, because the renderer
folds continuation lines before parsing inline spans.

But the folding does NOT apply to a **list item** whose bold spans a line break. The first
authored draft of 1.2.0 had exactly one such bullet, and the renderer emitted the opening
`**` verbatim in one paragraph and the rest in another. The source-level check flagged the
line, but 36 other flagged lines were harmless — so a source-level count cannot tell you
whether the DOCX is clean. Only opening the DOCX can.

Hence: this opens the artefact and fails on any literal marker.

Usage:  python tools/check_render.py [file.docx ...]
        (defaults to dist/PACSP-M-1.2.0.docx)
"""
from __future__ import annotations

import sys
from pathlib import Path

M = Path(__file__).resolve().parent.parent
MARKERS = ("**", "__")


def check(path: Path) -> tuple[int, list[str]]:
    try:
        from docx import Document
    except ImportError:
        print("  python-docx is not available; cannot inspect the DOCX.")
        print("  Use the bundled runtime, e.g. the one load_workspace_dependencies reports.")
        return -1, []
    doc = Document(str(path))

    leaks: list[str] = []
    for p in doc.paragraphs:
        if any(m in p.text for m in MARKERS):
            leaks.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if any(m in p.text for m in MARKERS):
                        leaks.append("[table] " + p.text)
    return len(leaks), leaks


def main() -> int:
    args = sys.argv[1:] or ["dist/PACSP-M-1.2.0.docx"]
    failed = 0
    for a in args:
        p = Path(a) if Path(a).is_absolute() else M / a
        if not p.exists():
            print(f"  [X] missing {p}")
            failed += 1
            continue
        n, leaks = check(p)
        if n < 0:
            return 2
        if n == 0:
            print(f"  [OK] {p.name}: inline markers paired, nothing leaked")
        else:
            failed += 1
            print(f"  [X]  {p.name}: {n} paragraph(s) carry a literal marker")
            for t in leaks[:6]:
                print(f"        {t[:96]!r}")
    print()
    print(f"  渲染标记自检：{len(args) - failed}/{len(args)} 通过")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
